import pandas as pd, numpy as np
from openpyxl import load_workbook
U="/mnt/user-data/uploads/v4_session_membership_estimates_wtf_.xlsx"; O="/mnt/user-data/outputs/bbg_dataset"
wb=load_workbook(U,read_only=True,data_only=True)
def rows(n): 
    ws=wb[n]; return [[c.value for c in r] for r in ws.iter_rows()]
# ---- membership ----
mem=[]
for idx in ["HSI","HSCEI","HSTECH"]:
    d=rows("IndexHist_"+idx)
    for c0 in range(0,len(d[1]),3):
        dt=d[1][c0]
        if dt is None: continue
        for r in d[2:]:
            if c0<len(r) and isinstance(r[c0],str) and r[c0].strip():
                w=r[c0+1] if c0+1<len(r) and isinstance(r[c0+1],(int,float)) else np.nan
                mem.append((idx,str(dt),r[c0]+" Equity",w))
M=pd.DataFrame(mem,columns=["index","asof","ticker","weight"]).drop_duplicates()
M["asof"]=pd.to_datetime(M["asof"],format="%Y%m%d")
M["weight"]=np.where(M["weight"].abs()<1e-6,np.nan,M["weight"])   # the -2.4e-14 placeholder = no weight returned
M.to_parquet(f"{O}/index_membership_history.parquet",index=False); M.to_csv(f"{O}/index_membership_history.csv",index=False)
print("membership:",M.shape)
print(M.groupby(["index"]).agg(quarters=("asof","nunique"),names=("ticker","nunique"),first=("asof","min"),last=("asof","max")).to_string())
print("\nHSI count by quarter (sample):")
print(M[M["index"]=="HSI"].groupby("asof").size().iloc[::6].to_string())
# ---- estimates ----
def pairs(sheet,col):
    d=rows(sheet); hdr=d[0]; out=[]
    for c in range(0,len(hdr),2):
        t=hdr[c]
        if not isinstance(t,str) or not t.strip(): continue
        for r in d[2:]:
            if c>=len(r): break
            dt=r[c]; v=r[c+1] if c+1<len(r) else None
            if hasattr(dt,"year") and isinstance(v,(int,float)): out.append((t,pd.Timestamp(dt),v))
    return pd.DataFrame(out,columns=["ticker","date",col])
e1=pairs("Est_EPS_FY1","eps_fy1"); e2=pairs("Est_EPS_FY2","eps_fy2"); tp=pairs("TargetPx","target_px")
EST=e1.set_index(["ticker","date"]).join([e2.set_index(["ticker","date"]),tp.set_index(["ticker","date"])],how="outer").reset_index()
EST.to_parquet(f"{O}/estimates_monthly.parquet",index=False)
print("\nestimates:",EST.shape,"tickers",EST.ticker.nunique(),"range",EST.date.min().date(),EST.date.max().date())
print(EST.notna().sum().to_string())
