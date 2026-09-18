import pandas as pd, numpy as np
from openpyxl import load_workbook
O="/mnt/user-data/outputs/bbg_dataset"
U="/mnt/user-data/uploads/v5_hsci_batch_saved_.xlsx"
wb=load_workbook(U,read_only=True,data_only=True)
def pairs(sheet,col):
    ws=wb[sheet]; d=[[c.value for c in r] for r in ws.iter_rows()]
    hdr=d[0]; out=[]
    for c in range(0,len(hdr),2):
        t=hdr[c]
        if not isinstance(t,str) or not t.strip(): continue
        for r in d[2:]:
            if c>=len(r): break
            dt=r[c]; v=r[c+1] if c+1<len(r) else None
            if hasattr(dt,"year"):
                out.append((t,pd.Timestamp(dt),v if isinstance(v,(int,float)) else np.nan))
    return pd.DataFrame(out,columns=["ticker","date",col])
fr=[]
for sh,col in [("Close","close"),("Volume","volume"),("Turnover","turnover"),("TotalReturnIdx","tri"),("MktCap","mktcap")]:
    fr.append(pairs(sh,col).set_index(["ticker","date"])[col])
new=pd.concat(fr,axis=1).reset_index()
print("batch1:",new.shape,"tickers",new.ticker.nunique(),new.date.min().date(),new.date.max().date())
old=pd.read_parquet(f"{O}/prices_daily.parquet")
old["turnover"]=np.nan
comb=pd.concat([old,new]).drop_duplicates(["ticker","date"],keep="last").sort_values(["ticker","date"])
comb.to_parquet(f"{O}/prices_daily_v2.parquet",index=False)
print("combined:",comb.shape,"tickers",comb.ticker.nunique())
# static
ws=wb["Static"]; d=[[c.value for c in r] for r in ws.iter_rows()]
sdf=pd.DataFrame(d[1:],columns=d[0]).rename(columns={"Ticker":"ticker"})
sdf=sdf[sdf.ticker.notna()&sdf.ticker.astype(str).str.contains("Equity")]
so=pd.read_csv(f"{O}/static_snapshot.csv")
s2=pd.concat([so,sdf]).drop_duplicates("ticker",keep="last")
s2.to_csv(f"{O}/static_snapshot_v2.csv",index=False); print("static:",s2.shape)
