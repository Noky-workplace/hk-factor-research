import pandas as pd, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as L
F=Font(name="Arial"); B=Font(name="Arial",bold=True); Y=PatternFill("solid",fgColor="FFFF00")
S,E="20160101","20260918"
MACRO="""Sub SafeFreeze()
    Dim ws As Worksheet, c As Range, bad As Long
    For Each ws In ThisWorkbook.Worksheets
        For Each c In ws.UsedRange
            If VarType(c.Value) = vbString Then
                If InStr(c.Value, "Requesting") > 0 Then bad = bad + 1
            End If
        Next c
    Next ws
    If bad > 0 Then
        MsgBox "STOP: " & bad & " cells still loading. Wait, then run again.", vbCritical
        Exit Sub
    End If
    For Each ws In ThisWorkbook.Worksheets
        ws.UsedRange.Value = ws.UsedRange.Value
    Next ws
    MsgBox "Frozen. Now Save As a new name.", vbInformation
End Sub"""
M=pd.read_parquet("/mnt/user-data/outputs/bbg_dataset/index_membership_history.parquet")
px=pd.read_parquet("/mnt/user-data/outputs/bbg_dataset/prices_daily_v2.parquet")
miss=[t for t in sorted(set(M["ticker"])-set(px.ticker.unique())) if t[0].isdigit()]
print(len(miss),"missing tickers")
wb=Workbook(); s=wb.active; s.title="Setup"
for r in [("PURPOSE","Session 6: (1) HSCI membership history - the fix that matters most  (2) prices for delisted/removed index members",""),
          ("WHY","Without these, every backtest holds only stocks that SURVIVED into today's index. Bias was ~2/3 of the measured return.",""),
          ("ORDER","Refresh Workbook -> check HSCI_Hist blocks differ between dates -> SafeFreeze -> Save As -> upload",""),
          ("IF STALLS","Delete the Delisted_* sheets, refresh, save; then reopen the original and do Delisted alone in a second pass.",""),
          ("NOTE","'#N/A Field Not Applicable' on early dates is normal and fine (index or stock did not exist yet).","")]: s.append(r)
s.column_dimensions["A"].width=12; s.column_dimensions["B"].width=125
mac=wb.create_sheet("MacroVBA")
for i,l in enumerate(MACRO.split("\n")): mac.cell(1+i,1,l)
mac.column_dimensions["A"].width=80
dates=[]; y,mo=2016,3
while (y,mo)<=(2026,9):
    d=(dt.date(y,mo,1)+dt.timedelta(days=31)).replace(day=1)-dt.timedelta(days=1); dates.append(d.strftime("%Y%m%d")); mo+=3
    if mo>12: mo=3; y+=1
w=wb.create_sheet("HSCI_Hist")
w["A1"]="Hang Seng COMPOSITE constituents + weights at each quarter-end. THE key sheet: check that blocks differ."
for k,d in enumerate(dates):
    c=k*3+1; cell=w.cell(2,c,d); cell.number_format="@"; cell.font=B
    w.cell(3,c,f'=BDS("HSCI Index","INDX_MWEIGHT_HIST","END_DATE_OVERRIDE","{d}")')
w.freeze_panes="A3"
# delisted / removed names: prices in chunks of 40 per sheet
CH=40
for bi in range(0,len(miss),CH):
    chunk=miss[bi:bi+CH]; p=wb.create_sheet(f"Delisted_{bi//CH+1}")
    for i,t in enumerate(chunk):
        c=2*i+1
        p.cell(1,c,t); p.cell(1,c).font=B; p.cell(2,c,"Date"); p.cell(2,c+1,"TRI")
        p.cell(3,c,f'=BDH("{t}","TOT_RETURN_INDEX_GROSS_DVDS","{S}","{E}")')
    p.freeze_panes="A3"
    p2=wb.create_sheet(f"DelistedPx_{bi//CH+1}")
    for i,t in enumerate(chunk):
        c=2*i+1
        p2.cell(1,c,t); p2.cell(1,c).font=B; p2.cell(2,c,"Date"); p2.cell(2,c+1,"PX_LAST")
        p2.cell(3,c,f'=BDH("{t}","PX_LAST","{S}","{E}","CshAdjNormal","Y","CshAdjAbnormal","Y","CapChg","Y")')
    p2.freeze_panes="A3"
lst=wb.create_sheet("DelistedList")
lst["A1"]="Tickers pulled on the Delisted_* sheets (historical index members with no price data yet)"; lst["A1"].font=B
for i,t in enumerate(miss): lst.cell(2+i,1,t)
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font.bold is not True: c.font=F
wb.save("/mnt/user-data/outputs/v6_hsci_history_delisted.xlsx"); print("ok",wb.sheetnames)
