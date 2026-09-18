from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as L
import datetime as dt
F=Font(name="Arial"); B=Font(name="Arial",bold=True); Y=PatternFill("solid",fgColor="FFFF00")
S,E="20160101","20260918"
MACRO="""Sub SafeFreeze()
    Dim ws As Worksheet, c As Range, bad As Long
    For Each ws In ThisWorkbook.Worksheets
        For Each c In ws.UsedRange
            If VarType(c.Value) = vbString Then
                If InStr(c.Value, "Requesting") > 0 Then bad = bad + 1
            End If
        Next c
    Next ws
    If bad > 0 Then
        MsgBox "STOP: " & bad & " cells still loading. Wait, then run again.", vbCritical
        Exit Sub
    End If
    For Each ws In ThisWorkbook.Worksheets
        ws.UsedRange.Value = ws.UsedRange.Value
    Next ws
    MsgBox "Frozen. Now Save As a new name.", vbInformation
End Sub"""
wb=Workbook(); s=wb.active; s.title="Setup"
for r in [("PURPOSE","Session 4: historical index membership (fixes survivorship bias), consensus estimates, entitlement tests",""),
          ("Order","1) Refresh Workbook  2) check IndexHist sheets show names+weights that DIFFER between dates  3) SafeFreeze  4) Save As + upload",""),
          ("IF #N/A","On EntitlementTest, #N/A Authorization = not entitled: ignore. Elsewhere, delete the offending sheet and refresh again.",""),
          ("NOTE","Estimates sheets: ~95 stocks x 3 fields = 285 BDH requests (monthly). Fine under the daily limit.","")]: s.append(r)
s.column_dimensions["A"].width=12; s.column_dimensions["B"].width=120
mac=wb.create_sheet("MacroVBA")
for i,l in enumerate(MACRO.split("\n")): mac.cell(1+i,1,l)
m=wb.create_sheet("Members"); m["A1"]='=BDS("HSI Index","INDX_MEMBERS")'
for r in range(1,301): m.cell(r,2,f'=IF(A{r}="","",A{r}&" Equity")')
# quarter-end dates as TEXT
dates=[]; y,mo=2016,3
while (y,mo)<=(2026,9):
    d=(dt.date(y,mo,1)+dt.timedelta(days=31)).replace(day=1)-dt.timedelta(days=1); dates.append(d.strftime("%Y%m%d")); mo+=3
    if mo>12: mo=3; y+=1
for idx in ["HSI Index","HSCEI Index","HSTECH Index"]:
    w=wb.create_sheet("IndexHist_"+idx.split()[0])
    w["A1"]=f"{idx}: constituents + weights at each quarter-end. Row 2 = date (TEXT). Each block spills below."
    for k,d in enumerate(dates):
        c=k*3+1; cell=w.cell(2,c,d); cell.number_format="@"; cell.font=B
        w.cell(3,c,f'=BDS("{idx}","INDX_MWEIGHT_HIST","END_DATE_OVERRIDE","{d}")')
    w.freeze_panes="A3"
def tkr(i): return f'IF(Members!$B${i}="","",Members!$B${i})'
for nm,f,ov in [("Est_EPS_FY1","BEST_EPS",',"BEST_FPERIOD_OVERRIDE","1FY"'),("Est_EPS_FY2","BEST_EPS",',"BEST_FPERIOD_OVERRIDE","2FY"'),
                ("TargetPx","BEST_TARGET_PRICE",'')]:
    p=wb.create_sheet(nm)
    for i in range(1,101):
        c=2*i-1; p.cell(1,c,f'={tkr(i)}'); p.cell(1,c).font=B; p.cell(2,c,"Date"); p.cell(2,c+1,f)
        p.cell(3,c,f'=IF({L(c)}$1="","",BDH({L(c)}$1,"{f}","{S}","{E}","Per","M"{ov}))')
    p.freeze_panes="A3"
fy=wb.create_sheet("FiscalYearEnd")
fy["A1"]="Ticker"; fy["B1"]="BEST_FPERIOD_END_DT (FY1)"; fy["C1"]="FISCAL_YEAR_PERIOD"
for i in range(1,101):
    r=i+1; fy.cell(r,1,f'={tkr(i)}')
    fy.cell(r,2,f'=IF($A{r}="","",BDP($A{r},"BEST_FPERIOD_END_DT","BEST_FPERIOD_OVERRIDE","1FY"))')
    fy.cell(r,3,f'=IF($A{r}="","",BDP($A{r},"FISCAL_YEAR_PERIOD"))')
t=wb.create_sheet("EntitlementTest")
t.append(["Field","Test on 700 HK Equity","Result meaning"])
for f,note in [("30DAY_IMPVOL_100.0MNY_DF","30d ATM implied vol"),("SHORT_INT","short interest"),("SHORT_INT_RATIO","days to cover"),
               ("ESG_DISCLOSURE_SCORE","ESG"),("NEWS_SENTIMENT_DAILY_AVG","news sentiment"),("PX_LAST","control - must work")]:
    t.append([f,f'=BDP("700 HK Equity","{f}")',note])
t.column_dimensions["A"].width=32; t.column_dimensions["B"].width=28; t.column_dimensions["C"].width=30
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font.bold is not True: c.font=F
wb.save("/mnt/user-data/outputs/v4_session_membership_estimates.xlsx"); print("ok", wb.sheetnames)
