import pandas as pd, numpy as np, os, json
from openpyxl import load_workbook
U="/mnt/user-data/uploads"; O="/mnt/user-data/outputs/bbg_dataset"; os.makedirs(O,exist_ok=True)

def rows(f,n):
    wb=load_workbook(f"{U}/{f}",read_only=True,data_only=True); ws=wb[n]
    r=[[c.value for c in row] for row in ws.iter_rows()]; wb.close(); return r

def num(v):
    return v if isinstance(v,(int,float)) and not isinstance(v,bool) else np.nan

def pairs_long(f,sheet,value_name,hdr_row=0,first_data=2,label_from="ticker"):
    """Sheets laid out as [ticker | blank | ticker | blank ...] with Date/Value pairs beneath."""
    r=rows(f,sheet); hdr=r[hdr_row]; out=[]
    for c in range(0,len(hdr),2):
        t=hdr[c]
        if not isinstance(t,str) or not t.strip(): continue
        for row in r[first_data:]:
            if c>=len(row): break
            d=row[c]; v=row[c+1] if c+1<len(row) else None
            if not hasattr(d,"year"): continue
            out.append((t,pd.Timestamp(d),num(v)))
    df=pd.DataFrame(out,columns=[label_from,"date",value_name])
    return df

manifest={}
# ---------- prices ----------
P="v3_2_prices_saved_.xlsx"
frames=[]
for sh,col in [("Close","close"),("Volume","volume"),("TotalReturnIdx","tri"),("MktCap","mktcap")]:
    d=pairs_long(P,sh,col); frames.append(d.set_index(["ticker","date"])[col])
prices=pd.concat(frames,axis=1).reset_index().sort_values(["ticker","date"])
prices=prices.drop_duplicates(["ticker","date"])
prices.to_parquet(f"{O}/prices_daily.parquet",index=False); prices.to_csv(f"{O}/prices_daily.csv",index=False)
manifest["prices_daily"]=dict(rows=len(prices),tickers=prices.ticker.nunique(),
    start=str(prices.date.min().date()),end=str(prices.date.max().date()),cols=list(prices.columns))

# ---------- fundamentals ----------
Fn="v3_3_fundamentals_saved_.xlsx"
fr=[]
for sh,col in [("Revenue","revenue"),("NetIncome","net_income"),("EPS","eps"),("BVPS","bvps"),
               ("CFO","cfo"),("TotDebt","total_debt"),("ROE","roe"),("SharesOut","shares_out")]:
    d=pairs_long(Fn,sh,col); fr.append(d.set_index(["ticker","date"])[col])
fund=pd.concat(fr,axis=1).reset_index().rename(columns={"date":"period_end"}).sort_values(["ticker","period_end"])
# Bloomberg sometimes emits both 2016-12-30 and 2016-12-31 rows; keep the one with data
fund["ym"]=fund.period_end.dt.to_period("M")
fund=fund.sort_values(["ticker","ym","period_end"]).groupby(["ticker","ym"]).last().reset_index().drop(columns="ym")
# attach announcement dates from salvaged earnings table
ed=pd.read_csv("/mnt/user-data/outputs/bbg_salvaged/earnings_dates.csv",parse_dates=["announce_date"])
def per_end(fp):
    try:
        y,p=fp.split(":")
        y=int(y)
        return {"S1":pd.Timestamp(y,6,30),"A":pd.Timestamp(y,12,31),"S2":pd.Timestamp(y,12,31),
                "Q1":pd.Timestamp(y,3,31),"Q2":pd.Timestamp(y,6,30),"Q3":pd.Timestamp(y,9,30),"Q4":pd.Timestamp(y,12,31)}.get(p)
    except Exception: return None
ed["period_end"]=ed.fiscal_period.map(per_end)
ed=ed.dropna(subset=["period_end"]).sort_values("announce_date").drop_duplicates(["ticker","period_end"])
fund["pm"]=fund.period_end.dt.to_period("M"); ed["pm"]=ed.period_end.dt.to_period("M")
fund=fund.merge(ed[["ticker","pm","announce_date"]],on=["ticker","pm"],how="left").drop(columns="pm")
# fallback: if no announcement date, assume 90 days after period end (conservative)
fund["announce_date_est"]=fund.announce_date.isna()
fund["announce_date"]=fund.announce_date.fillna(fund.period_end+pd.Timedelta(days=90))
fund.to_parquet(f"{O}/fundamentals_semiannual.parquet",index=False); fund.to_csv(f"{O}/fundamentals_semiannual.csv",index=False)
manifest["fundamentals_semiannual"]=dict(rows=len(fund),tickers=fund.ticker.nunique(),
    announce_matched=int((~fund.announce_date_est).sum()),announce_estimated=int(fund.announce_date_est.sum()),cols=list(fund.columns))

# ---------- macro / benchmarks ----------
M="v3_1_macro_static_saved_.xlsx"
mf=[]
for sh in ["Rates_Credit","FX","Commodities","GlobalEquity","MacroData"]:
    d=pairs_long(M,sh,"value",label_from="series")
    d["group"]=sh; mf.append(d)
macro=pd.concat(mf).dropna(subset=["value"]).drop_duplicates(["series","date"])
macro.to_parquet(f"{O}/macro_series.parquet",index=False); macro.to_csv(f"{O}/macro_series.csv",index=False)
wide=macro.pivot(index="date",columns="series",values="value").sort_index()
wide.to_csv(f"{O}/macro_wide.csv")
manifest["macro_series"]=dict(rows=len(macro),series=sorted(macro.series.unique().tolist()))

# ---------- static ----------
r=rows(M,"Static"); hdr=r[0]
st=pd.DataFrame([row for row in r[1:] if isinstance(row[0],str) and row[0].strip()],columns=hdr)
for c in st.columns[1:]:
    if c not in ("NAME","GICS_SECTOR_NAME","ID_ISIN"): st[c]=st[c].map(num)
st=st.rename(columns={"Ticker":"ticker"})
st.to_csv(f"{O}/static_snapshot.csv",index=False)
manifest["static_snapshot"]=dict(rows=len(st),cols=list(st.columns))

# copy salvaged
import shutil
for f in ["earnings_dates.csv","dividends.csv","members_HSI.csv"]:
    shutil.copy(f"/mnt/user-data/outputs/bbg_salvaged/{f}",f"{O}/{f}")
json.dump(manifest,open(f"{O}/MANIFEST.json","w"),indent=2,default=str)
print(json.dumps(manifest,indent=2,default=str))
